"""A deliberately minimal Excel export of the monthly report: just the
title of each completed task and its estimated hours — nothing else
(no financials, no actual hours, no categories). For handing to a
client/manager who only wants "what was done and how many hours it was
budgeted for", not the internal profitability numbers.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def render_monthly_report_xlsx(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش ماهانه"
    ws.sheet_view.rightToLeft = True

    ws.append([f"{report['project_name']} — {report['period_start']:%Y-%m} تا {report['period_end']:%Y-%m-%d}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(["عنوان کار", "ساعت تخمینی"])
    for col in (1, 2):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="right" if col == 1 else "center")

    for task in report["completed_tasks"]:
        ws.append([task["title"], float(task["estimated_hours"]) if task["estimated_hours"] is not None else None])
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(horizontal="center")

    ws.column_dimensions[get_column_letter(1)].width = 60
    ws.column_dimensions[get_column_letter(2)].width = 16

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
